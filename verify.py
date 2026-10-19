"""
Independent checks on the finished workbook.

1. Recalculate the workbook in LibreOffice (a different calculation engine from the one that wrote
   the cached values) and compare every formula result with the value stored in the file.
2. Confirm the file contains no Excel error values.
3. Compare the cleaned orders with the ground truth the raw exports were generated from.
"""
import os
import shutil
import subprocess
import time

import openpyxl

ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!")


RECALC_ALWAYS = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>
</oor:items>
"""


def recalculate(xlsx_path, out_dir):
    """Recalculate in headless LibreOffice and return (path of the recalculated copy, seconds).

    LibreOffice normally trusts the cached values of files that look Excel-made, so the profile
    is set to ALWAYS recalculate on load. Otherwise this check would compare a file with itself."""
    os.makedirs(out_dir, exist_ok=True)
    home = os.path.join(out_dir, "_lohome")
    profile = os.path.join(home, ".config", "libreoffice", "4", "user")
    os.makedirs(profile, exist_ok=True)
    with open(os.path.join(profile, "registrymodifications.xcu"), "w", encoding="utf-8") as f:
        f.write(RECALC_ALWAYS)
    env = {**os.environ, "HOME": home}
    t0 = time.time()
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", out_dir, xlsx_path],
                   check=True, timeout=240, env=env, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    shutil.rmtree(home, ignore_errors=True)
    return os.path.join(out_dir, os.path.basename(xlsx_path)), time.time() - t0


def tamper_copy(xlsx_path, out_path):
    """Copy the workbook with ONE cached formula result changed by +1. Used to prove the check bites."""
    import re
    import zipfile
    with zipfile.ZipFile(xlsx_path) as zin, zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED) as zout:
        done = False
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet2.xml" and not done:  # the By Region sheet
                text = data.decode("utf-8")
                m = re.search(r"(<f>[^<]*</f><v>)([0-9.]+)(</v>)", text)
                text = text[:m.start()] + m.group(1) + str(float(m.group(2)) + 1) + m.group(3) + text[m.end():]
                data, done = text.encode("utf-8"), True
            zout.writestr(item, data)


def compare_formulas(original, recalculated):
    """Compare cached values in `original` with LibreOffice's recalculated values."""
    wf = openpyxl.load_workbook(original)                    # formulas
    wc = openpyxl.load_workbook(original, data_only=True)    # cached results
    wr = openpyxl.load_workbook(recalculated, data_only=True)  # recalculated by LibreOffice
    checked = mismatches = errors = 0
    max_dev, per_sheet = 0.0, {}
    for ws in wf.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type != "f":
                    continue
                cached = wc[ws.title][cell.coordinate].value
                fresh = wr[ws.title][cell.coordinate].value
                checked += 1
                per_sheet[ws.title] = per_sheet.get(ws.title, 0) + 1
                if isinstance(fresh, str) and fresh in ERRORS or isinstance(cached, str) and cached in ERRORS:
                    errors += 1
                if isinstance(cached, (int, float)) and isinstance(fresh, (int, float)) and not isinstance(cached, bool):
                    dev = abs(cached - fresh)
                    max_dev = max(max_dev, dev)
                    if dev > 1e-6:
                        mismatches += 1
                elif cached != fresh:
                    mismatches += 1
    return {"formulas": checked, "mismatches": mismatches, "errors": errors, "max_abs_deviation": max_dev,
            "by_sheet": per_sheet}


def compare_with_truth(records, truth):
    """records: cleaned orders; truth: the orders the exports were generated from."""
    t = {x["order_id"]: x for x in truth}
    r = {x["order_id"]: x for x in records}
    same_ids = set(t) == set(r)
    wrong = sum(1 for k in r if k in t and (r[k]["amount_c"] != t[k]["amount_c"] or r[k]["region"] != t[k]["region"]
                                           or r[k]["order_date"] != t[k]["order_date"] or r[k]["units"] != t[k]["units"]))
    return {"true_orders": len(t), "cleaned_orders": len(r), "same_order_ids": same_ids, "orders_with_wrong_field": wrong,
            "revenue_matches": sum(x["amount_c"] for x in t.values()) == sum(x["amount_c"] for x in r.values())}
